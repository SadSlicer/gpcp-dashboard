"""GPCP Dashboard — professional Excel (.xlsx) portfolio report.

`build_excel_report(static, price_history, snapshot, settings)` returns the
workbook as bytes, built entirely in memory (nothing touches the server's disk
— the container is shared by every user). One sheet per topic, mirroring the
app — Summary, Holdings, Transactions, Realized P&L, Performance, then every
Analytics sub-tab: Risk Metrics, Calendar (monthly returns), Attribution,
Correlation, Benchmark, Monte Carlo, Exposure — plus Prices and Notes.

Every figure is read off the same series and formulas as the app (unit value
base 100 at the first close, the Risk Metrics EAR / Sharpe / Sortino
conventions, the heatmap's month-end logic, the Benchmark tab's common window,
the Monte Carlo engine), using the user's current Analytics settings
(risk-free rate, benchmark, attribution window, Monte Carlo inputs).

Sort / filter safety: the only formulas are ROW-LOCAL (P&L = value − cost on
the same row) or use ABSOLUTE ranges (weight = value / SUM($I$5:$I$9)), so
sorting a table keeps them right. Totals sit BELOW a blank spacer row — outside
the filter range and outside Excel's "current region" — and use SUBTOTAL(109,…)
so they follow the active filter. Order-dependent series (daily return,
drawdown) are stored as values, never as formulas.

openpyxl only (pure Python — no native wheel, nothing that can segfault on
Streamlit Cloud).
"""
from __future__ import annotations

import datetime as dt
import io
import math
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines, DateAxis
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (CharacterProperties, Font as DFont, Paragraph,
                                   ParagraphProperties, RegularTextRun)
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import data
import pro

# ---- house style (the PDF report's palette: navy ink, zinc neutrals) -------
NAVY = "0F2547"
INK = "1E2C3F"
GREY = "64748B"
AXIS = "94A3B8"
GRIDC = "E2E8F0"
RULE = "CDD6E2"
ZEBRA = "F5F7FA"
SOFT = "EEF2F7"
GOLD = "B08D57"
POS = "0F766E"
NEG = "B3261E"
POS_SOFT = "7FB9A6"
NEG_SOFT = "D98E7A"
SLICES = ["0F2547", "0F766E", "B08D57", "5B7FA6", "7A5C9E", "B0705A",
          "3E7C59", "8A5A7A", "2F4A6E", "A3B4C8", "C9A66B", "4F9C8C"]

FONT = "Calibri"
F_TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=10, color=GREY)
F_SECTION = Font(name=FONT, size=11, bold=True, color=NAVY)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BODY = Font(name=FONT, size=10, color=INK)
F_BOLD = Font(name=FONT, size=10, bold=True, color=INK)
F_LABEL = Font(name=FONT, size=10, color=GREY)
F_NOTE = Font(name=FONT, size=9, italic=True, color=GREY)

FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_ZEBRA = PatternFill("solid", fgColor=ZEBRA)
FILL_TOTAL = PatternFill("solid", fgColor=SOFT)

THIN = Side(style="thin", color=RULE)
B_BOTTOM = Border(bottom=THIN)
B_TOTAL = Border(top=Side(style="thin", color=NAVY), bottom=Side(style="double", color=NAVY))

FMT_PCT = '0.00%;[Red]-0.00%'
FMT_PCT_SIGNED = '+0.00%;[Red]-0.00%;0.00%'
FMT_PTS = '+0.000" pts";[Red]-0.000" pts";0.000" pts"'
FMT_NUM = '#,##0.00;[Red]-#,##0.00'
FMT_NUM4 = '#,##0.0000'
FMT_SHARES = '#,##0.####'
FMT_DATE = 'dd mmm yyyy'
FMT_RATIO = '0.00;[Red]-0.00'
FMT_INT = '#,##0'

CHART_W, CHART_H = 26.0, 11.5     # cm


def _money_fmt(sym: str, decimals: int = 2) -> str:
    z = "0." + "0" * decimals if decimals else "0"
    return f'#,##0{z[1:]} "{sym}";[Red]-#,##0{z[1:]} "{sym}"'


def _clean(v):
    """None for NaN / inf / missing — Excel shows an empty cell, not 'nan'."""
    if v is None:
        return None
    if isinstance(v, (float, np.floating)):
        f = float(v)
        return None if (math.isnan(f) or math.isinf(f)) else f
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime().date()
    return v


def safe_filename(name: str) -> str:
    """Portfolio name → a filesystem-safe stem ("My PEA (EUR)" → "My_PEA_EUR")."""
    s = re.sub(r"[^\w\-]+", "_", (name or "").strip(), flags=re.UNICODE).strip("_")
    return s or "Portfolio"


# ---------------------------------------------------------------------------
# Chart styling
# ---------------------------------------------------------------------------

def _cp(size: int, color: str, bold: bool = False) -> CharacterProperties:
    return CharacterProperties(sz=size, b=bold, solidFill=color,
                               latin=DFont(typeface=FONT))


def _rich(size: int, color: str, bold: bool = False) -> RichText:
    cp = _cp(size, color, bold)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[], endParaRPr=cp)])


def _title(text: str, size: int = 1200, color: str = NAVY) -> Title:
    cp = _cp(size, color, True)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def _style_axis(ax, title: str | None = None, fmt: str | None = None,
                grid: bool = False) -> None:
    ax.delete = False                       # openpyxl ≥3.1 hides axes otherwise
    ax.txPr = _rich(900, GREY)
    ax.spPr = GraphicalProperties(ln=LineProperties(solidFill=AXIS, w=9525))
    ax.majorTickMark = "out"
    if title:
        ax.title = _title(title, 900, GREY)
    if fmt:
        ax.numFmt = NumFmt(formatCode=fmt, sourceLinked=False)
    ax.majorGridlines = (ChartLines(spPr=GraphicalProperties(
        ln=LineProperties(solidFill=GRIDC, w=6350))) if grid else None)


def _finish(chart, title: str, *, legend: bool = True, w: float = CHART_W,
            h: float = CHART_H) -> None:
    chart.title = _title(title)
    chart.width, chart.height = w, h
    if legend:
        chart.legend.position = "b"
        chart.legend.txPr = _rich(900, INK)
    else:
        chart.legend = None
    chart.graphical_properties = GraphicalProperties(
        ln=LineProperties(solidFill=RULE, w=9525))
    chart.plot_area.graphicalProperties = GraphicalProperties(noFill=True)


def _line_series(s, color: str, width_pt: float = 1.75, dash: str | None = None) -> None:
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = int(width_pt * 12700)
    if dash:
        s.graphicalProperties.line.dashStyle = dash
    s.marker.symbol = "none"
    s.smooth = False


def _date_line_chart(ws, top: int, n: int, x_col: int, y_cols: list[int], title: str,
                     y_title: str, y_fmt: str, colors: list[str],
                     dashes: list[str | None] | None = None,
                     widths: list[float] | None = None, legend: bool = True,
                     span_days: int | None = None) -> LineChart:
    """Line chart over a DATE axis (real time scale, month ticks)."""
    ch = LineChart()
    ch.x_axis = DateAxis(crossAx=100)
    ch.x_axis.axPos = "b"
    ch.y_axis.crossAx = 500
    for j, col in enumerate(y_cols):
        ch.add_data(Reference(ws, min_col=col, min_row=top, max_row=top + n),
                    titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=x_col, min_row=top + 1, max_row=top + n))
    for j, s in enumerate(ch.series):
        _line_series(s, colors[j % len(colors)],
                     (widths or [1.75] * len(ch.series))[j],
                     (dashes or [None] * len(ch.series))[j])
    months = max(1, (span_days or 365) // 30)
    ch.x_axis.majorTimeUnit = "months"
    ch.x_axis.majorUnit = max(1, round(months / 10))
    ch.x_axis.baseTimeUnit = "days"
    _style_axis(ch.x_axis, None, "mmm yy")
    ch.x_axis.tickLblPos = "low"
    _style_axis(ch.y_axis, y_title, y_fmt, grid=True)
    _finish(ch, title, legend=legend)
    return ch


def _date_area_chart(ws, top: int, n: int, x_col: int, y_col: int, title: str,
                     y_title: str, y_fmt: str, color: str,
                     span_days: int | None = None) -> AreaChart:
    ch = AreaChart()
    ch.x_axis = DateAxis(crossAx=100)
    ch.x_axis.axPos = "b"
    ch.y_axis.crossAx = 500
    ch.add_data(Reference(ws, min_col=y_col, min_row=top, max_row=top + n),
                titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=x_col, min_row=top + 1, max_row=top + n))
    s = ch.series[0]
    s.graphicalProperties.solidFill = color
    s.graphicalProperties.line.solidFill = color
    months = max(1, (span_days or 365) // 30)
    ch.x_axis.majorTimeUnit = "months"
    ch.x_axis.majorUnit = max(1, round(months / 10))
    _style_axis(ch.x_axis, None, "mmm yy")
    ch.x_axis.tickLblPos = "low"
    _style_axis(ch.y_axis, y_title, y_fmt, grid=True)
    _finish(ch, title, legend=False)
    return ch


def _bar_chart(ws, top: int, n: int, x_col: int, y_col: int, title: str,
               y_title: str, y_fmt: str, *, horizontal: bool = False,
               sign_colors: bool = True, values: list | None = None,
               color: str = NAVY, x_title: str | None = None,
               labels: bool = False, label_fmt: str | None = None,
               w: float = CHART_W, h: float = CHART_H) -> BarChart:
    """Bar chart; bars tinted green / red by sign when `sign_colors`."""
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.gapWidth = 40
    ch.add_data(Reference(ws, min_col=y_col, min_row=top, max_row=top + n),
                titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=x_col, min_row=top + 1, max_row=top + n))
    s = ch.series[0]
    s.graphicalProperties.solidFill = color
    s.graphicalProperties.line.noFill = True
    if sign_colors and values is not None:
        for i, v in enumerate(values):
            if v is None:
                continue
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = POS if v >= 0 else NEG
            pt.graphicalProperties.line.noFill = True
            s.dPt.append(pt)
    if labels:
        s.dLbls = DataLabelList()
        s.dLbls.showVal = True
        s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
        s.dLbls.txPr = _rich(800, INK)
        if label_fmt:
            s.dLbls.numFmt = label_fmt
    _style_axis(ch.x_axis, x_title)
    ch.x_axis.tickLblPos = "low"
    _style_axis(ch.y_axis, y_title, y_fmt, grid=True)
    _finish(ch, title, legend=False, w=w, h=h)
    return ch


def _doughnut(ws, top: int, n: int, label_col: int, value_col: int, title: str,
              w: float = 13.5, h: float = 10.5) -> DoughnutChart:
    """Doughnut with one brand colour per slice and % labels."""
    ch = DoughnutChart()
    ch.holeSize = 55
    ch.add_data(Reference(ws, min_col=value_col, min_row=top, max_row=top + n),
                titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=label_col, min_row=top + 1, max_row=top + n))
    s = ch.series[0]
    for i in range(n):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = SLICES[i % len(SLICES)]
        pt.graphicalProperties.line.solidFill = "FFFFFF"
        pt.graphicalProperties.line.width = 12700
        s.dPt.append(pt)
    s.dLbls = DataLabelList()
    s.dLbls.showPercent = True
    s.dLbls.showVal = s.dLbls.showSerName = s.dLbls.showCatName = False
    s.dLbls.showLegendKey = False
    s.dLbls.numFmt = "0.0%"
    s.dLbls.txPr = _rich(900, "FFFFFF", True)
    _finish(ch, title, w=w, h=h)
    ch.legend.position = "r"
    return ch


# ---------------------------------------------------------------------------
# Metrics — same conventions as the app
# ---------------------------------------------------------------------------

def _period_returns(vl: pd.Series) -> dict[str, float | None]:
    """MTD / YTD / 12M / since inception off the unit value (as the Holdings
    TOTAL row): reference = last close at or before the window start, the
    first close (= 100) for a window that starts before inception."""
    out = {"MTD": None, "YTD": None, "12M": None, "SI": None}
    if vl is None or len(vl) < 1:
        return out
    last_d, last_v = vl.index[-1], float(vl.iloc[-1])
    starts = {
        "MTD": last_d.normalize().replace(day=1),
        "YTD": last_d.normalize().replace(month=1, day=1),
        "12M": last_d.normalize() - pd.DateOffset(years=1),
    }
    for k, start in starts.items():
        before = vl[vl.index <= start]
        base = float(before.iloc[-1]) if len(before) else float(vl.iloc[0])
        out[k] = (last_v / base - 1.0) if base > 0 else None
    first = float(vl.iloc[0])
    out["SI"] = (last_v / first - 1.0) if first > 0 else None
    return out


def _risk_stats(vl: pd.Series, rf: float) -> dict:
    """The Risk Metrics tab's formulas, plus the drawdown episode details."""
    keys = ("n", "mean_d", "std_d", "ann_return", "ann_vol", "sharpe", "sortino",
            "max_dd", "var95", "var99", "cvar95", "calmar", "peak", "trough",
            "recovery", "dd_days")
    out: dict = dict.fromkeys(keys)
    if vl is None or len(vl) < 3:
        return out
    td = pro.TRADING_DAYS
    rets = vl.pct_change().dropna()
    n = len(rets)
    mean_d, std_d = float(rets.mean()), float(rets.std(ddof=1))
    ann_return = (1.0 + mean_d) ** td - 1.0
    ann_vol = std_d * math.sqrt(td) if std_d > 0 else 0.0
    rf_d = (1.0 + rf) ** (1.0 / td) - 1.0
    excess = rets - rf_d
    ann_excess = (1.0 + excess.mean()) ** td - 1.0
    below = np.minimum(excess.values, 0.0)
    ann_down = math.sqrt(float(np.mean(below ** 2))) * math.sqrt(td)
    dd = vl / vl.cummax() - 1.0
    max_dd = float(dd.min())
    trough = dd.idxmin()
    if max_dd < 0:
        peak = vl.loc[:trough].idxmax()
        after = vl.loc[trough:]
        rec = after[after >= vl.loc[peak]]
        recovery = rec.index[0] if not rec.empty else None
        dd_days = ((recovery or vl.index[-1]) - peak).days
    else:
        peak, recovery, dd_days = trough, None, 0
    var95 = float(np.quantile(rets, 0.05))
    tail = rets[rets <= var95]
    out.update(
        n=n, mean_d=mean_d, std_d=std_d, ann_return=ann_return, ann_vol=ann_vol,
        sharpe=(ann_excess / ann_vol) if ann_vol > 0 else None,
        sortino=(ann_excess / ann_down) if ann_down > 0 else None,
        max_dd=max_dd, var95=var95, var99=float(np.quantile(rets, 0.01)),
        cvar95=float(tail.mean()) if len(tail) else var95,
        calmar=(ann_return / abs(max_dd)) if abs(max_dd) > 1e-9 else None,
        peak=peak, trough=trough, recovery=recovery, dd_days=dd_days,
    )
    return out


def _monthly_matrix(vl: pd.Series) -> tuple[pd.DataFrame, pd.Series]:
    """(year × month matrix + 'Year' column, chronological monthly returns).

    Calendar Heatmap rules: a month needs both real month-ends; the year column
    is close-to-close (first close for the first year) → current year = YTD.
    """
    cols = list(range(1, 13)) + ["Year"]
    if vl is None or len(vl) < 2:
        return pd.DataFrame(columns=cols), pd.Series(dtype=float)
    m_ret = vl.resample("ME").last().pct_change().dropna()
    years = sorted({int(y) for y in vl.index.year}, reverse=True)
    mat = pd.DataFrame(index=years, columns=cols, dtype=float)
    for ts, r in m_ret.items():
        mat.at[ts.year, ts.month] = float(r)
    for y in years:
        in_year = vl[vl.index.year == y]
        prev = vl[vl.index.year == (y - 1)]
        base = float(prev.iloc[-1]) if len(prev) else float(in_year.iloc[0])
        if base > 0:
            mat.at[y, "Year"] = float(in_year.iloc[-1]) / base - 1.0
    return mat, m_ret


def _to_days(s: pd.Series) -> pd.Series:
    idx = pd.DatetimeIndex(s.index)
    if idx.tz is not None:
        idx = idx.tz_localize(None)
    s = s.copy()
    s.index = idx.normalize()
    return s[~s.index.duplicated(keep="last")]


def _benchmark(vl: pd.Series, choice: str | None, custom: str | None):
    """Benchmark tab logic: (label, portfolio idx, benchmark idx) on the common
    window, both rebased to 100 — or None when unavailable. The benchmark is
    converted into the portfolio currency (label says e.g. "(USD → EUR)")."""
    if vl is None or len(vl) < 2 or not pro.prices.YAHOO_ENABLED:
        return None
    if choice in pro.BENCHMARKS:
        info = pro.BENCHMARKS[choice]
    elif (custom or "").strip():
        info = {"ticker": custom.strip(), "label": custom.strip().upper()}
    else:
        info = next(iter(pro.BENCHMARKS.values()))
    s, e = vl.index[0].date(), vl.index[-1].date()
    pf_ccy = data.current_portfolio_currency()
    fx_note = pro._benchmark_fx_label(info["ticker"] or "IWDA.AS", pf_ccy)
    label = f"{info['label']} ({fx_note})" if fx_note else info["label"]
    if info["ticker"] is None:
        world = pro._fetch_benchmark("IWDA.AS", s, e, pf_ccy)
        bond = pro._fetch_benchmark("AGGH.AS", s, e, pf_ccy)
        if world.empty or bond.empty:
            return None
        j = pd.concat([world, bond], axis=1, keys=["w", "b"]).ffill().dropna()
        j = j / j.iloc[0] * 100
        bm = 0.6 * j["w"] + 0.4 * j["b"]
    else:
        bm = pro._fetch_benchmark(info["ticker"], s, e, pf_ccy)
    if bm is None or bm.empty:
        return None
    nav, bm = _to_days(vl), _to_days(bm)
    common = nav.index.intersection(bm.index).sort_values()
    if len(common) < 2:
        return None
    nav_a, bm_a = nav.loc[common], bm.loc[common]
    return label, nav_a / nav_a.iloc[0] * 100, bm_a / bm_a.iloc[0] * 100


# ---------------------------------------------------------------------------
# Sheet helpers
# ---------------------------------------------------------------------------

class _Ctx:
    def __init__(self, pf_name: str, subtitle: str, sym: str):
        self.pf_name, self.subtitle, self.sym = pf_name, subtitle, sym


def _sheet(wb, ctx: _Ctx, name: str, title: str, extra: str = "",
           width_cols: int = 8, first: bool = False):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = ctx.subtitle + (f"  ·  {extra}" if extra else "")
    ws["A2"].font = F_SUB
    for c in range(1, width_cols + 1):
        ws.cell(row=2, column=c).border = Border(bottom=Side(style="medium", color=NAVY))
    ws.row_dimensions[1].height = 24
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.left.text = f"{ctx.pf_name} — Portfolio report"
    ws.oddFooter.right.text = "Page &P / &N"
    return ws


def _section(ws, row: int, col: int, text: str, span: int = 2) -> None:
    c = ws.cell(row=row, column=col, value=text)
    c.font = F_SECTION
    for k in range(col, col + span):
        ws.cell(row=row, column=k).border = Border(bottom=Side(style="thin", color=NAVY))


def _table(ws, top: int, headers: list[str], rows: list[list], fmts: list[str | None],
           *, left: int = 1, total: list | None = None, autofilter: bool = True,
           zebra: bool = True) -> dict:
    """Header + body (+ total row BELOW a blank spacer row).

    Returns {"first", "last", "total", "next"}: first/last data rows, the total
    row (or None) and the next free row. The header is top-aligned and tall so
    the filter drop-down (bottom-right of the cell) never covers its text.
    """
    for j, h in enumerate(headers):
        c = ws.cell(row=top, column=left + j, value=h)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="left" if fmts[j] is None else "right",
                                vertical="top", wrap_text=True)
    ws.row_dimensions[top].height = 34
    first = top + 1
    last = top + len(rows)
    for i, row in enumerate(rows):
        r = first + i
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=left + j, value=_clean(v))
            c.font = F_BODY
            c.border = B_BOTTOM
            if fmts[j]:
                c.number_format = fmts[j]
                c.alignment = Alignment(horizontal="right")
            if zebra and i % 2 == 1:
                c.fill = FILL_ZEBRA
    if autofilter and rows:
        ws.auto_filter.ref = (f"{get_column_letter(left)}{top}:"
                              f"{get_column_letter(left + len(headers) - 1)}{last}")
    total_row = None
    nxt = last + 2
    if total is not None and rows:
        total_row = last + 2          # blank spacer row at last + 1
        for j, v in enumerate(total):
            c = ws.cell(row=total_row, column=left + j, value=_clean(v))
            c.font, c.fill, c.border = F_BOLD, FILL_TOTAL, B_TOTAL
            if fmts[j]:
                c.number_format = fmts[j]
                c.alignment = Alignment(horizontal="right")
        nxt = total_row + 2
    return {"first": first, "last": last, "total": total_row, "next": nxt}


def _kv(ws, top: int, title: str, items: list[tuple], *, col: int = 1,
        note_col: bool = False) -> int:
    """Section + label / value (/ note) rows. Returns the next free row."""
    _section(ws, top, col, title, 3 if note_col else 2)
    r = top
    for i, it in enumerate(items):
        label, value, fmt = it[0], it[1], it[2]
        r = top + 1 + i
        lc = ws.cell(row=r, column=col, value=label)
        lc.font, lc.border = F_LABEL, B_BOTTOM
        v = _clean(value)
        vc = ws.cell(row=r, column=col + 1, value=v if v is not None else "—")
        vc.font, vc.border = F_BOLD, B_BOTTOM
        vc.alignment = Alignment(horizontal="right")
        if fmt and v is not None:
            vc.number_format = fmt
        if note_col:
            nc = ws.cell(row=r, column=col + 2, value=it[3] if len(it) > 3 else None)
            nc.font, nc.border = F_NOTE, B_BOTTOM
            nc.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[r].height = 28
    return r + 2


def _note(ws, row: int, text: str, col: int = 1, span: int = 8) -> None:
    c = ws.cell(row=row, column=col, value=text)
    c.font = F_NOTE
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=col, end_row=row + 1,
                   end_column=col + span - 1)


def _widths(ws, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _sub(rng: str) -> str:
    return f"SUBTOTAL(109,{rng})"


# ---------------------------------------------------------------------------
# Main builder
# ---------------------------------------------------------------------------

def build_excel_report(static, price_history: pd.DataFrame, snapshot: dict,
                       settings: dict | None = None) -> bytes:
    settings = settings or {}
    pf = data.current_portfolio() or {}
    pf_name = (pf.get("name") or "").strip() or "Portfolio"
    pf_ccy = data.current_portfolio_currency()
    sym = data.CURRENCY_SYMBOL.get(pf_ccy, pf_ccy)
    money = _money_fmt(sym)
    money0 = _money_fmt(sym, 0)
    today = dt.date.today()
    positions = snapshot.get("positions", {}) or {}
    inception = snapshot.get("inception_date")
    rf = float(settings.get("rf", pro.DEFAULT_RF))

    vlf = data.compute_vl_series(price_history)
    if vlf is not None and not vlf.empty:
        vlf = vlf.copy()
        vlf["date"] = pd.to_datetime(vlf["date"])
        vlf = vlf.sort_values("date").reset_index(drop=True)
        vl = pd.Series(vlf["vl"].astype(float).values, index=vlf["date"])
    else:
        vlf = pd.DataFrame(columns=["date", "nav", "units", "vl", "net_invested"])
        vl = pd.Series(dtype=float)
    as_of = vl.index[-1].date() if len(vl) else today
    span = (vl.index[-1] - vl.index[0]).days if len(vl) > 1 else 365

    periods = _period_returns(vl)
    risk = _risk_stats(vl, rf)
    held = sorted(((a, p) for a, p in positions.items() if a != "Cash"),
                  key=lambda kv: -(kv[1].get("value") or 0))

    wb = Workbook()
    wb.properties.title = f"{pf_name} — Portfolio report"
    wb.properties.creator = pf_name
    wb.properties.subject = "Portfolio holdings, performance, risk and analytics"
    ctx = _Ctx(pf_name, f"{pf_name}  ·  Base currency {pf_ccy}  ·  Data as of "
                        f"{as_of:%d %b %Y}  ·  Generated {today:%d %b %Y}", sym)

    # ===================== Summary =====================
    ws = _sheet(wb, ctx, "Summary", f"{pf_name} — Portfolio Report", width_cols=12,
                first=True)
    r = _kv(ws, 4, "Portfolio", [
        ("Portfolio", pf_name, None),
        ("Base currency", pf_ccy, None),
        ("Inception date", inception, FMT_DATE),
        ("Data as of", as_of, FMT_DATE),
        ("Lines held", len(held), "0"),
    ])
    r = _kv(ws, r, "Key figures", [
        ("Net asset value", snapshot.get("total_value"), money),
        ("Net invested (total cost)", snapshot.get("net_invested"), money),
        ("Profit / loss", snapshot.get("cash_pnl_eur"), money),
        ("Total return (money-weighted)", snapshot.get("total_return_pct"), FMT_PCT_SIGNED),
        ("Unit value (base 100)", snapshot.get("vl"), FMT_NUM),
        ("Cash balance", snapshot.get("cash_balance"), money),
        ("Daily P&L", snapshot.get("daily_pnl_eur"), money),
        ("Daily change", snapshot.get("daily_pnl_pct"), FMT_PCT_SIGNED),
    ])
    r = _kv(ws, r, "Performance (time-weighted, unit value)", [
        ("Month to date", periods["MTD"], FMT_PCT_SIGNED),
        ("Year to date", periods["YTD"], FMT_PCT_SIGNED),
        ("12 months", periods["12M"], FMT_PCT_SIGNED),
        ("Since inception", periods["SI"], FMT_PCT_SIGNED),
        ("Annualized return (EAR)", risk["ann_return"], FMT_PCT_SIGNED),
    ])
    r = _kv(ws, r, f"Risk (risk-free {rf:.2%})", [
        ("Annualized volatility", risk["ann_vol"], FMT_PCT),
        ("Sharpe ratio", risk["sharpe"], FMT_RATIO),
        ("Sortino ratio", risk["sortino"], FMT_RATIO),
        ("Maximum drawdown", risk["max_dd"], FMT_PCT_SIGNED),
        ("VaR 95% (1 day)", risk["var95"], FMT_PCT_SIGNED),
    ])
    _note(ws, r, "For information only — not investment advice. Prices: Yahoo Finance "
                 "adjusted close (dividends reinvested). Figures may differ from a "
                 "custodian statement. Past performance does not guarantee future results.",
          span=2)
    _widths(ws, {"A": 34, "B": 20, "C": 3})
    # Summary charts are drawn once the Performance / Holdings data exist (below).
    ws_summary = ws

    # ===================== Holdings =====================
    ws = _sheet(wb, ctx, "Holdings", "Holdings", width_cols=13)
    rows = []
    for a, p in held:
        cost_pf = None
        if p.get("inception_price") and p.get("shares"):
            cost_pf = p["inception_price"] * p["shares"] * (p.get("fx_rate_cost") or 1.0)
        rows.append([a, data.TICKER_BY_ASSET.get(a, ""), data.ISIN_BY_ASSET.get(a, ""),
                     p.get("currency"), p.get("shares"), p.get("price"),
                     p.get("inception_price"), cost_pf, p.get("value"), None, None, None,
                     p.get("daily_return_pf", p.get("daily_return"))])
    cash = positions.get("Cash")
    if cash:
        rows.append(["Cash", "", "", pf_ccy, None, None, None, cash.get("value"),
                     cash.get("value"), None, None, None, None])
    top = 4
    f, l = top + 1, top + len(rows)
    for i, row in enumerate(rows):      # row-local + absolute formulas (sort-safe)
        rr = f + i
        row[9] = f"=I{rr}-H{rr}"
        row[10] = f'=IF(H{rr}=0,"",J{rr}/H{rr})'
        row[11] = f"=I{rr}/SUM($I${f}:$I${l})"
    tot = l + 2
    t = _table(ws, top, ["Asset", "Ticker", "ISIN", "Ccy", "Shares", "Last price (native)",
                         "Avg cost (native)", f"Cost ({sym})", f"Market value ({sym})",
                         f"P&L ({sym})", "Return", "Weight", "Daily change"], rows,
               [None, None, None, None, FMT_SHARES, FMT_NUM, FMT_NUM, money, money, money,
                FMT_PCT_SIGNED, FMT_PCT, FMT_PCT_SIGNED],
               total=["TOTAL (visible rows)", "", "", "", None, None, None,
                      "=" + _sub(f"H{f}:H{l}"), "=" + _sub(f"I{f}:I{l}"),
                      "=" + _sub(f"J{f}:J{l}"), f'=IF(H{tot}=0,"",J{tot}/H{tot})',
                      "=" + _sub(f"L{f}:L{l}"), None])
    _note(ws, t["next"], "Totals follow the active filter (SUBTOTAL). P&L, Return and "
                         "Weight are formulas — they stay correct when the table is sorted "
                         "or filtered.", span=10)
    if rows:
        ws.add_chart(_doughnut(ws, top, len(rows), 1, 9, "Portfolio weights",
                               w=18, h=11), f"A{t['next'] + 3}")
    _widths(ws, {"A": 38, "B": 11, "C": 15, "D": 6, "E": 11, "F": 13, "G": 13,
                 "H": 16, "I": 18, "J": 15, "K": 11, "L": 11, "M": 12})
    ws.freeze_panes = "B5"
    holdings_rows = (top, len(rows))

    # ===================== Transactions =====================
    ws = _sheet(wb, ctx, "Transactions", "Transactions", "amounts in trade currency")
    try:
        tx = data.load_transactions()
    except Exception:
        tx = pd.DataFrame()
    rows = []
    if tx is not None and not tx.empty:
        for _, x in tx.sort_values("Date").iterrows():
            rows.append([x.get("Date"), x.get("Type"), x.get("Asset"), x.get("ISIN"),
                         x.get("Shares"), x.get("Price"), x.get("Total"), x.get("Currency")])
    _table(ws, 4, ["Date", "Type", "Asset", "ISIN", "Shares", "Price", "Amount", "Ccy"],
           rows, [FMT_DATE, None, None, None, FMT_SHARES, FMT_NUM, FMT_NUM, None])
    _widths(ws, {"A": 14, "B": 10, "C": 38, "D": 15, "E": 12, "F": 12, "G": 14, "H": 7})
    ws.freeze_panes = "A5"

    # ===================== Realized P&L =====================
    try:
        sells = data.sell_pnl_rows()
    except Exception:
        sells = []
    if sells:
        ws = _sheet(wb, ctx, "Realized P&L", "Realized P&L", "vs weighted-average cost")
        rows = []
        for i, s in enumerate(sorted(sells, key=lambda s: s["date"])):
            rr = 5 + i
            rows.append([s.get("date"), s.get("asset"), s.get("shares"), s.get("sell_price"),
                         s.get("avg_cost"), f'=IF(E{rr}="","",(D{rr}-E{rr})*C{rr})',
                         f'=IF(OR(E{rr}="",E{rr}=0),"",D{rr}/E{rr}-1)', s.get("currency")])
        _table(ws, 4, ["Date", "Asset", "Shares sold", "Sale price", "Avg cost",
                       "Realized P&L", "Return", "Ccy"], rows,
               [FMT_DATE, None, FMT_SHARES, FMT_NUM, FMT_NUM, FMT_NUM, FMT_PCT_SIGNED, None])
        _widths(ws, {"A": 14, "B": 38, "C": 12, "D": 12, "E": 12, "F": 14, "G": 11, "H": 7})
        ws.freeze_panes = "A5"

    # ===================== Performance (daily) =====================
    ws = _sheet(wb, ctx, "Performance", "Daily performance",
                "unit value = time-weighted, base 100 at the first close")
    rows = []
    if not vlf.empty:
        dd = vlf["vl"] / vlf["vl"].cummax() - 1.0
        dret = vlf["vl"].pct_change()
        for i, row in vlf.iterrows():
            rr = 5 + i
            rows.append([row["date"], row.get("nav"), row.get("net_invested"),
                         f"=B{rr}-C{rr}", row.get("units"), row.get("vl"),
                         dret.iat[i], dd.iat[i]])
    _table(ws, 4, ["Date", f"NAV ({sym})", f"Net invested ({sym})", f"P&L ({sym})",
                   "Units", "Unit value", "Daily return", "Drawdown"], rows,
           [FMT_DATE, money, money, money, FMT_NUM4, FMT_NUM, FMT_PCT_SIGNED,
            FMT_PCT_SIGNED])
    n_perf = len(rows)
    if n_perf >= 2:
        ws.add_chart(_date_line_chart(ws, 4, n_perf, 1, [6], "Unit value (base 100)",
                                      "Unit value", "0", [NAVY], widths=[2.0],
                                      legend=False, span_days=span), "J4")
        ws.add_chart(_date_line_chart(ws, 4, n_perf, 1, [2, 3],
                                      f"Net asset value vs net invested ({sym})",
                                      f"Amount ({sym})", "#,##0", [NAVY, GOLD],
                                      dashes=[None, "dash"], widths=[2.0, 1.5],
                                      span_days=span), "J28")
        ws.add_chart(_date_area_chart(ws, 4, n_perf, 1, 8, "Drawdown from peak",
                                      "Drawdown", "0%", NEG_SOFT, span_days=span), "J52")
    _widths(ws, {"A": 14, "B": 16, "C": 18, "D": 15, "E": 13, "F": 12, "G": 13, "H": 12})
    ws.freeze_panes = "B5"

    # Summary charts (reference the Performance / Holdings sheets)
    if n_perf >= 2:
        ws_summary.add_chart(_date_line_chart(ws, 4, n_perf, 1, [6],
                                              "Unit value (base 100)", "Unit value", "0",
                                              [NAVY], widths=[2.0], legend=False,
                                              span_days=span), "E4")
    if holdings_rows[1]:
        hws = wb["Holdings"]
        ws_summary.add_chart(_doughnut(hws, holdings_rows[0], holdings_rows[1], 1, 9,
                                       "Portfolio weights", w=16, h=10.5), "E28")

    # ===================== Risk Metrics =====================
    ws = _sheet(wb, ctx, "Risk Metrics", "Risk metrics",
                f"daily unit-value returns · risk-free {rf:.2%}", width_cols=4)
    peak, trough, recov = risk["peak"], risk["trough"], risk["recovery"]
    r = _kv(ws, 4, "Return & risk", [
        ("Annualized return (EAR)", risk["ann_return"], FMT_PCT_SIGNED,
         "<0 loss · 0-7% modest · 7-12% good · >12% very good"),
        ("Mean daily return", risk["mean_d"], '+0.0000%;[Red]-0.0000%', ""),
        ("Annualized volatility", risk["ann_vol"], FMT_PCT,
         "<10% low · 10-20% moderate · >20% high"),
        ("Daily volatility (σ)", risk["std_d"], '0.0000%', ""),
        ("Sharpe ratio", risk["sharpe"], FMT_RATIO,
         "<0 poor · 0-1 average · 1-2 good · >2 excellent"),
        ("Sortino ratio", risk["sortino"], FMT_RATIO,
         "Downside-only risk · >1 good · >2 very good"),
        ("Calmar ratio", risk["calmar"], FMT_RATIO,
         "Return / |max DD| · <1 weak · 1-3 good · >3 excellent"),
    ], note_col=True)
    r = _kv(ws, r, "Drawdown", [
        ("Maximum drawdown", risk["max_dd"], FMT_PCT_SIGNED,
         ">-10% comfortable · -10 to -30% normal for equities · <-40% severe"),
        ("Peak date", peak, FMT_DATE, "Last high before the worst fall"),
        ("Trough date", trough, FMT_DATE, "Bottom of the worst fall"),
        ("Recovery date", recov if recov is not None else "Not recovered",
         FMT_DATE if recov is not None else None, "First close back at the peak"),
        ("Drawdown duration (days)", risk["dd_days"], FMT_INT, "Peak → recovery (or today)"),
    ], note_col=True)
    r = _kv(ws, r, "Tail risk (1 day, historical)", [
        ("VaR 95%", risk["var95"], FMT_PCT_SIGNED, "Loss exceeded ~1 day in 20"),
        ("VaR 99%", risk["var99"], FMT_PCT_SIGNED, "Loss exceeded ~1 day in 100"),
        ("CVaR 95% (expected shortfall)", risk["cvar95"], FMT_PCT_SIGNED,
         "Average loss beyond the VaR 95%"),
        ("Daily returns observed", risk["n"], FMT_INT,
         "Annualized figures are noisy below ~60 days"),
    ], note_col=True)
    # Histogram of daily returns (bins as a table → column chart)
    rets = vl.pct_change().dropna() if len(vl) > 2 else pd.Series(dtype=float)
    if len(rets) >= 5:
        counts, edges = np.histogram(rets.values, bins=max(20, min(40, int(math.sqrt(len(rets))))))
        _section(ws, r, 1, "Distribution of daily returns", 3)
        hrows = [[f"{(edges[i] + edges[i + 1]) / 2 * 100:+.2f}%", (edges[i] + edges[i + 1]) / 2,
                  int(counts[i])] for i in range(len(counts))]
        ht = _table(ws, r + 1, ["Bin (mid)", "Return", "Days"], hrows,
                    [None, FMT_PCT_SIGNED, FMT_INT], autofilter=False)
        ch = _bar_chart(ws, r + 1, len(hrows), 1, 3, "Distribution of daily returns",
                        "Number of days", "0", sign_colors=True,
                        values=[x[1] for x in hrows], x_title="Daily return")
        ch.gapWidth = 10
        ws.add_chart(ch, "F30")
    if n_perf >= 2:
        pws = wb["Performance"]
        ws.add_chart(_date_area_chart(pws, 4, n_perf, 1, 8, "Drawdown curve", "Drawdown",
                                      "0%", NEG_SOFT, span_days=span), "F4")
    _widths(ws, {"A": 32, "B": 16, "C": 52, "D": 3})

    # ===================== Calendar (monthly returns) =====================
    ws = _sheet(wb, ctx, "Monthly Returns", "Calendar — monthly returns",
                "unit value, month-end to month-end", width_cols=14)
    mat, m_ret = _monthly_matrix(vl)
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
              "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    rows = [[int(y)] + [mat.at[y, m] for m in range(1, 13)] + [mat.at[y, "Year"]]
            for y in mat.index]
    t = _table(ws, 4, ["Year"] + months + ["Year / YTD"], rows,
               ["0"] + [FMT_PCT_SIGNED] * 13, autofilter=False, zebra=False)
    if rows:
        ws.conditional_formatting.add(f"B5:M{4 + len(rows)}", ColorScaleRule(
            start_type="num", start_value=-0.05, start_color="F4C7C3",
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=0.05, end_color="B7E1CD"))
        for i in range(len(rows)):
            ws.cell(row=5 + i, column=1).alignment = Alignment(horizontal="left")
            ws.cell(row=5 + i, column=14).font = F_BOLD
    _note(ws, t["next"], "A month is shown once both its own and the previous month-end "
                         "close exist. The Year column is close-to-close (the current year "
                         "is year-to-date).", span=14)
    r = t["next"] + 3
    # Yearly returns (oldest → newest) + chronological monthly returns
    yrows = [[str(int(y)), mat.at[y, "Year"]] for y in sorted(mat.index)]
    if yrows:
        _section(ws, r, 1, "Calendar-year returns", 2)
        yt = _table(ws, r + 1, ["Year", "Return"], yrows, [None, FMT_PCT_SIGNED],
                    autofilter=False)
        ws.add_chart(_bar_chart(ws, r + 1, len(yrows), 1, 2, "Calendar-year returns",
                                "Return", "0%", values=[x[1] for x in yrows],
                                labels=True, label_fmt="+0.0%;-0.0%", w=16, h=10),
                     f"D{r}")
        r = max(yt["next"], r + 22)
    if len(m_ret):
        mrows = [[ts.strftime("%b %y"), float(v)] for ts, v in m_ret.items()]
        _section(ws, r, 1, "Monthly returns (chronological)", 2)
        _table(ws, r + 1, ["Month", "Return"], mrows, [None, FMT_PCT_SIGNED],
               autofilter=False)
        ws.add_chart(_bar_chart(ws, r + 1, len(mrows), 1, 2, "Monthly returns",
                                "Return", "0%", values=[x[1] for x in mrows]), f"D{r}")
    _widths(ws, {"A": 10, **{get_column_letter(c): 9.5 for c in range(2, 14)}, "N": 12})
    ws.freeze_panes = "B5"

    # ===================== Performance Attribution =====================
    _attribution_sheet(wb, ctx, static, price_history, settings, money)

    # ===================== Correlation =====================
    _correlation_sheet(wb, ctx, price_history, positions)

    # ===================== Benchmark =====================
    ws = _sheet(wb, ctx, "Benchmark", "Benchmark comparison",
                "both indexed to 100 on the first common date", width_cols=6)
    bmr = None
    try:
        bmr = _benchmark(vl, settings.get("bm_choice"), settings.get("bm_custom"))
    except Exception:
        bmr = None
    if bmr is None:
        ws.cell(row=4, column=1, value="Benchmark data unavailable (no common history "
                                       "or price provider offline).").font = F_NOTE
    else:
        label, p_idx, b_idx = bmr
        pr, br = p_idx.pct_change().dropna(), b_idx.pct_change().dropna()
        ci = pr.index.intersection(br.index)
        pr, br = pr.loc[ci], br.loc[ci]
        ex = pr - br
        td = pro.TRADING_DAYS
        alpha = ex.mean() * td if len(ex) else None
        te = ex.std(ddof=1) * math.sqrt(td) if len(ex) > 1 else None
        beta = (float(np.cov(pr, br, ddof=1)[0, 1] / br.var(ddof=1))
                if len(ci) >= 2 and br.var() > 0 else None)
        p_tot, b_tot = p_idx.iloc[-1] / 100 - 1, b_idx.iloc[-1] / 100 - 1
        r = _kv(ws, 4, f"{pf_name} vs {label}", [
            ("Window", f"{p_idx.index[0]:%d %b %Y} → {p_idx.index[-1]:%d %b %Y}", None),
            ("Common trading days", len(p_idx), FMT_INT),
            (f"{pf_name} return", p_tot, FMT_PCT_SIGNED),
            (f"{label} return", b_tot, FMT_PCT_SIGNED),
            ("Outperformance", p_tot - b_tot, FMT_PCT_SIGNED),
            ("Alpha (annualized)", alpha, FMT_PCT_SIGNED),
            ("Tracking error", te, FMT_PCT),
            ("Information ratio", (alpha / te) if (alpha is not None and te) else None,
             FMT_RATIO),
            ("Beta", beta, FMT_RATIO),
            ("Correlation", float(pr.corr(br)) if len(ci) >= 2 else None, FMT_RATIO),
        ])
        top = r
        rows = []
        for i, d in enumerate(p_idx.index):
            rr = top + 1 + i
            rows.append([d, float(p_idx.loc[d]), float(b_idx.loc[d]), f"=B{rr}-C{rr}"])
        _table(ws, top, ["Date", pf_name, label, "Gap (index pts)"], rows,
               [FMT_DATE, FMT_NUM, FMT_NUM, '+0.00;[Red]-0.00'])
        n = len(rows)
        ws.add_chart(_date_line_chart(ws, top, n, 1, [2, 3],
                                      f"{pf_name} vs {label} (base 100)", "Index",
                                      "0", [NAVY, GOLD], dashes=[None, "sysDash"],
                                      widths=[2.0, 1.75], span_days=span), "F4")
        ws.add_chart(_date_area_chart(ws, top, n, 1, 4, "Cumulative outperformance",
                                      "Index points", "0",
                                      POS_SOFT if (p_tot - b_tot) >= 0 else NEG_SOFT,
                                      span_days=span), "F28")
        ws.freeze_panes = f"A{top + 1}"
    _widths(ws, {"A": 26, "B": 22, "C": 16, "D": 15})

    # ===================== Monte Carlo =====================
    _monte_carlo_sheet(wb, ctx, vl, snapshot, settings, money0)

    # ===================== Exposure =====================
    ws = _sheet(wb, ctx, "Exposure", "Look-through exposure",
                "value-weighted, invested assets only", width_cols=7)
    try:
        geo_agg, sec_agg, inv_total = pro._lookthrough_exposure(positions)
    except Exception:
        geo_agg, sec_agg, inv_total = {}, {}, 1.0
    chart_row = 4
    for col, title, agg in ((1, "Geographic", geo_agg), (5, "Sector", sec_agg)):
        _section(ws, 4, col, title, 3)
        items = sorted(agg.items(), key=lambda kv: -kv[1])
        if not items:
            ws.cell(row=5, column=col, value="No breakdown available.").font = F_NOTE
            continue
        L, V, W = (get_column_letter(col), get_column_letter(col + 2),
                   get_column_letter(col + 1))
        f, l = 6, 5 + len(items)
        rows = [[k, f"={V}{6 + i}/SUM(${V}${f}:${V}${l})", v]
                for i, (k, v) in enumerate(items)]
        t = _table(ws, 5, ["Bucket", "Weight", f"Value ({sym})"], rows,
                   [None, FMT_PCT, money], left=col, autofilter=False,
                   total=["Total (visible rows)", "=" + _sub(f"{W}{f}:{W}{l}"),
                          "=" + _sub(f"{V}{f}:{V}{l}")])
        chart_row = max(chart_row, t["next"] + 1)
    for col, title, agg in ((1, "Geographic exposure", geo_agg),
                            (5, "Sector exposure", sec_agg)):
        if agg:
            n = len(agg)
            ws.add_chart(_doughnut(ws, 5, n, col, col + 2, title, w=14, h=11),
                         f"{get_column_letter(col)}{chart_row}")
    _widths(ws, {"A": 26, "B": 10, "C": 16, "D": 4, "E": 26, "F": 10, "G": 16})

    # ===================== Prices =====================
    ws = _sheet(wb, ctx, "Prices", "Daily closing prices",
                f"adjusted close, converted to {pf_ccy}", width_cols=6)
    try:
        ph = data.price_history_in_portfolio_currency(price_history)
    except Exception:
        ph = price_history
    assets = [a for a, _ in held if a in getattr(ph, "columns", [])]
    rows = []
    if ph is not None and not ph.empty and assets:
        ph = ph.sort_values("date")
        if inception:
            ph = ph[pd.to_datetime(ph["date"]).dt.date >= inception]
        for _, row in ph.iterrows():
            rows.append([row["date"]] + [row.get(a) for a in assets])
    _table(ws, 4, ["Date"] + assets, rows, [FMT_DATE] + [FMT_NUM] * len(assets))
    for j in range(len(assets)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 20
    ws.column_dimensions["A"].width = 14
    ws.freeze_panes = "B5"

    # ===================== Notes =====================
    ws = _sheet(wb, ctx, "Notes", "Methodology", width_cols=2)
    notes = [
        ("Unit value", "Time-weighted performance: base 100 at the first day's close. Units "
                       "are created on deposits and destroyed on withdrawals at the prevailing "
                       "unit value, so cash flows never show up as performance."),
        ("Total return", "Money-weighted: (NAV − net invested) / net invested. Differs from "
                         "the unit value when capital is added at different dates."),
        ("MTD / YTD / 12M", "Unit value at the last date over the last close at or before the "
                            "window start; windows starting before inception read from 100."),
        ("Annualized return", "Effective annual rate on the mean daily return: (1 + r̄)^252 − 1."),
        ("Volatility", "Standard deviation of daily unit-value returns × √252."),
        ("Sharpe / Sortino", f"Annualized excess return over a {rf:.2%} risk-free rate, "
                             "divided by volatility / downside deviation."),
        ("VaR / CVaR", "Historical 1-day quantile of daily returns; CVaR = mean of the tail."),
        ("Attribution", "Contribution = start weight × asset return over the window."),
        ("Correlation", "Pairwise daily-return correlation over each pair's common dates; "
                        "diversification score = (1 − weight-weighted ρ̄) × 100."),
        ("Benchmark", "Converted into the portfolio currency at each day's FX rate "
                      "(unhedged), then both series rebased to 100 on the first common "
                      "date; alpha = mean excess × 252, tracking error = σ(excess) × √252."),
        ("Monte Carlo", "Geometric Brownian motion calibrated on the unit value's daily "
                        "returns, starting from today's NAV; contributions smoothed daily."),
        ("Prices", "Yahoo Finance adjusted close (splits and dividends), forward-filled over "
                   "foreign-market holidays, converted at the daily FX rate."),
        ("Exposure", "Issuer factsheet compositions, weighted by current market value."),
        ("Formulas", "Row formulas only reference their own row or absolute ranges, so "
                     "sorting and filtering keep them correct. Totals use SUBTOTAL and "
                     "follow the active filter."),
    ]
    for i, (k, v) in enumerate(notes):
        kc = ws.cell(row=4 + i, column=1, value=k)
        kc.font = F_BOLD
        kc.alignment = Alignment(vertical="top")
        vc = ws.cell(row=4 + i, column=2, value=v)
        vc.font = F_BODY
        vc.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[4 + i].height = 30
    _widths(ws, {"A": 22, "B": 110})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Analytics sheets
# ---------------------------------------------------------------------------

def _attribution_sheet(wb, ctx: _Ctx, static, price_history, settings, money) -> None:
    """Performance Attribution tab: start weight × return per asset."""
    sym = ctx.sym
    try:
        hist = (data.price_history_in_portfolio_currency(price_history)
                .sort_values("date").reset_index(drop=True))
        inc = data.get_inception_date()
        hist = hist[hist["date"].dt.date >= inc].reset_index(drop=True)
    except Exception:
        hist = pd.DataFrame()
    a_from = settings.get("att_from")
    a_to = settings.get("att_to")
    if not hist.empty:
        mn, mx = hist["date"].min().date(), hist["date"].max().date()
        a_from = a_from if (a_from and mn <= a_from <= mx) else mn
        a_to = a_to if (a_to and mn <= a_to <= mx and a_to > a_from) else mx
    ws = _sheet(wb, ctx, "Attribution", "Performance attribution",
                (f"{a_from:%d %b %Y} → {a_to:%d %b %Y}" if a_from else ""), width_cols=6)
    if hist.empty or not a_from or a_from >= a_to:
        ws.cell(row=4, column=1, value="Not enough history for an attribution.").font = F_NOTE
        return
    sub = hist[(hist["date"].dt.date >= a_from) & (hist["date"].dt.date <= a_to)]
    if len(sub) < 2:
        ws.cell(row=4, column=1, value="The window has fewer than 2 sessions.").font = F_NOTE
        return
    last = sub.iloc[-1]
    pf_ccy = data.current_portfolio_currency()
    cost_map = getattr(data, "avg_cost_by_asset", lambda: {})()
    per_start: dict[str, float] = {}
    for a in (pro._held_assets() or list(data.ASSETS)):
        if a not in sub.columns or pd.isna(last[a]):
            continue
        col_valid = sub[a].dropna()
        if col_valid.empty:
            continue
        cb = cost_map.get(a)
        if cb and a_from <= cb[0] <= a_to:
            a_ccy = (static.currencies.get(a) or pf_ccy).upper()
            rate = 1.0 if a_ccy == pf_ccy else data.fx_rate(a_ccy, pf_ccy, cb[0])
            per_start[a] = float(cb[1]) * rate
        else:
            per_start[a] = float(col_valid.iloc[0])
    items = []
    for a, p0 in per_start.items():
        sh = static.shares.get(a, 0)
        items.append((a, p0 * sh, float(last[a]) * sh))
    if not items:
        ws.cell(row=4, column=1, value="No valued asset over this window.").font = F_NOTE
        return
    nav0 = sum(x[1] for x in items)
    nav1 = sum(x[2] for x in items)
    contrib = {a: ((s0 / nav0) * (s1 / s0 - 1.0) if nav0 and s0 else 0.0)
               for a, s0, s1 in items}
    items.sort(key=lambda x: -contrib[x[0]])
    r = _kv(ws, 4, "Window", [
        ("From", a_from, FMT_DATE), ("To", a_to, FMT_DATE),
        (f"NAV start ({sym})", nav0, money), (f"NAV end ({sym})", nav1, money),
        ("Total return (period)", (nav1 / nav0 - 1.0) if nav0 else None, FMT_PCT_SIGNED),
    ])
    top = r
    f, l = top + 1, top + len(items)
    rows = []
    for i, (a, s0, s1) in enumerate(items):
        rr = f + i
        rows.append([a, s0, s1, f"=B{rr}/SUM($B${f}:$B${l})",
                     f'=IF(B{rr}=0,"",C{rr}/B{rr}-1)', f'=IF(B{rr}=0,"",D{rr}*E{rr}*100)'])
    tot = l + 2
    t = _table(ws, top, ["Asset", f"Start value ({sym})", f"End value ({sym})",
                         "Start weight", "Return", "Contribution"], rows,
               [None, money, money, FMT_PCT, FMT_PCT_SIGNED, FMT_PTS],
               total=["TOTAL (visible rows)", "=" + _sub(f"B{f}:B{l}"),
                      "=" + _sub(f"C{f}:C{l}"), "=" + _sub(f"D{f}:D{l}"),
                      f'=IF(B{tot}=0,"",C{tot}/B{tot}-1)', "=" + _sub(f"F{f}:F{l}")])
    _note(ws, t["next"], "Contribution (percentage points) = start weight × asset return. "
                         "Σ contributions ≈ total return (small gap = weight drift).", span=6)
    ws.add_chart(_bar_chart(ws, top, len(items), 1, 6, "Contribution to total return",
                            "Percentage points", "0.00", horizontal=True,
                            values=[contrib[a] for a, _, _ in items], labels=True,
                            label_fmt='+0.00;-0.00', h=max(8.0, 1.1 * len(items) + 4)),
                 f"A{t['next'] + 3}")
    ws.add_chart(_bar_chart(ws, top, len(items), 1, 5, "Asset return over the window",
                            "Return", "0%", values=[(s1 / s0 - 1.0) if s0 else 0.0
                                                    for _, s0, s1 in items],
                            labels=True, label_fmt="+0.0%;-0.0%", w=20, h=11), "H4")
    _widths(ws, {"A": 38, "B": 18, "C": 18, "D": 13, "E": 12, "F": 14})


def _correlation_sheet(wb, ctx: _Ctx, price_history, positions) -> None:
    """Correlation Matrix tab: pairwise ρ + diversification score."""
    ws = _sheet(wb, ctx, "Correlation", "Correlation matrix",
                "daily returns, pairwise over common dates", width_cols=6)
    held = pro._held_assets() or list(data.ASSETS)
    try:
        inc = data.get_inception_date()
        ph = price_history[price_history["date"].dt.date >= inc]
        df = ph.sort_values("date").set_index("date")[[a for a in held if a in ph.columns]]
        rets = df.pct_change()
    except Exception:
        rets = pd.DataFrame()
    if rets.empty or rets.dropna(how="all").empty or rets.shape[1] < 1:
        ws.cell(row=4, column=1, value="Not enough history for correlations.").font = F_NOTE
        return
    corr = rets.corr()
    valid = rets.notna().astype(int)
    n_obs = valid.T.dot(valid)
    names = list(corr.columns)
    k = len(names)
    hdr = ws.cell(row=4, column=1, value="ρ")
    hdr.font, hdr.fill = F_HEAD, FILL_HEAD
    for j, a in enumerate(names):
        c = ws.cell(row=4, column=2 + j, value=a)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
    ws.row_dimensions[4].height = 48
    for i, a in enumerate(names):
        c = ws.cell(row=5 + i, column=1, value=a)
        c.font, c.border = F_BOLD, B_BOTTOM
        for j, b in enumerate(names):
            v = _clean(corr.iloc[i, j])
            cc = ws.cell(row=5 + i, column=2 + j, value=v)
            cc.number_format = "0.00"
            cc.font = F_BODY
            cc.alignment = Alignment(horizontal="center")
            cc.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    rng = f"B5:{get_column_letter(1 + k)}{4 + k}"
    ws.conditional_formatting.add(rng, ColorScaleRule(
        start_type="num", start_value=-1, start_color="B3261E",
        mid_type="num", mid_value=0, mid_color="E8D9A8",
        end_type="num", end_value=1, end_color="0F766E"))
    r = 6 + k
    pairs = []
    for i, a in enumerate(names):
        for j, b in enumerate(names):
            if j > i and pd.notna(corr.iloc[i, j]):
                pairs.append((a, b, float(corr.iloc[i, j]), int(n_obs.iloc[i, j])))
    w = {a: float((positions.get(a, {}) or {}).get("allocation", 0.0) or 0.0) for a in names}
    num = sum(w[a] * w[b] * rho for a, b, rho, _ in pairs)
    den = sum(w[a] * w[b] for a, b, _, _ in pairs)
    avg = (num / den) if den > 0 else (np.mean([p[2] for p in pairs]) if pairs else None)
    score = (1 - avg) * 100 if avg is not None else None
    r = _kv(ws, r, "Diversification", [
        ("Weight-weighted average ρ", avg, "0.00"),
        ("Diversification score (/100)", score, "0.0"),
        ("Reading", ">60 good diversification · <30 highly correlated", None),
    ])
    if pairs:
        _section(ws, r, 1, "All pairs (sortable)", 4)
        prow = [[a, b, rho, n] for a, b, rho, n in sorted(pairs, key=lambda p: -p[2])]
        _table(ws, r + 1, ["Asset A", "Asset B", "Correlation", "Common days"], prow,
               [None, None, "0.000", FMT_INT])
    _widths(ws, {"A": 34, **{get_column_letter(2 + j): 16 for j in range(max(k, 3))}})


def _monte_carlo_sheet(wb, ctx: _Ctx, vl: pd.Series, snapshot: dict, settings: dict,
                       money0: str) -> None:
    """Monte Carlo tab with the user's current inputs."""
    sym = ctx.sym
    ws = _sheet(wb, ctx, "Monte Carlo", "Monte Carlo projection",
                "geometric Brownian motion calibrated on the unit value", width_cols=6)
    if vl is None or len(vl) < 3:
        ws.cell(row=4, column=1, value="Not enough history to calibrate.").font = F_NOTE
        return
    td = pro.TRADING_DAYS
    rets = vl.pct_change().dropna()
    mu_d, sigma_d = float(rets.mean()), float(rets.std(ddof=1))
    horizon = int(settings.get("mc_horizon") or 5)
    n_sims = int(settings.get("mc_n") or 10000)
    contrib_m = float(settings.get("mc_contrib") or 0)
    nav0 = float(snapshot.get("total_value") or 0.0) or float(vl.iloc[-1])
    target = float(settings.get("mc_target") or nav0 * 2)
    n_steps = horizon * td
    contrib_d = contrib_m * 12 / td
    p10, p50, p90, terminal, _ = pro._mc_simulate(nav0, mu_d, sigma_d, n_steps,
                                                   n_sims, contrib_d)
    inv_now = float(snapshot.get("net_invested") or 0.0) or nav0
    invested = inv_now + contrib_d * np.arange(1, n_steps + 1)
    inv_end = float(invested[-1])
    t10, t50, t90 = (float(np.percentile(terminal, 10)), float(np.median(terminal)),
                     float(np.percentile(terminal, 90)))
    r = _kv(ws, 4, "Inputs", [
        ("Starting NAV", nav0, money0),
        ("Horizon (years)", horizon, "0"),
        ("Simulations", n_sims, FMT_INT),
        (f"Monthly contribution ({sym})", contrib_m, money0),
        ("NAV target", target, money0),
        ("Calibrated return μ (annual)", mu_d * td, FMT_PCT_SIGNED),
        ("Calibrated volatility σ (annual)", sigma_d * math.sqrt(td), FMT_PCT),
    ])
    r = _kv(ws, r, f"Probabilities at {horizon} year(s)", [
        ("P(target reached)", float((terminal >= target).mean()) if target > 0 else None,
         "0.0%"),
        ("P(loss vs today's NAV)", float((terminal < nav0).mean()), "0.0%"),
    ])
    _section(ws, r, 1, f"At the horizon — {horizon} year(s)", 5)
    top = r + 1
    f = top + 1
    srows = []
    for i, (name, v) in enumerate((("Pessimistic (P10)", t10), ("Median (P50)", t50),
                                   ("Optimistic (P90)", t90))):
        rr = f + i
        srows.append([name, v, inv_end, f"=B{rr}-C{rr}", f'=IF(C{rr}=0,"",D{rr}/C{rr})'])
    t = _table(ws, top, ["Scenario", f"Projected NAV ({sym})", f"Total cost ({sym})",
                         f"Profit ({sym})", "Return on cost"], srows,
               [None, money0, money0, money0, FMT_PCT_SIGNED], autofilter=False)
    # Monthly path table (sampled every ~21 trading days) → fan chart
    dates = pd.bdate_range(vl.index[-1] + pd.Timedelta(days=1), periods=n_steps)
    step = 21
    idx = list(range(step - 1, n_steps, step))
    if idx[-1] != n_steps - 1:
        idx.append(n_steps - 1)
    r = t["next"] + 1
    _section(ws, r, 1, "Projected paths (monthly)", 5)
    prow = [[dates[i], float(p10[i]), float(p50[i]), float(p90[i]), float(invested[i])]
            for i in idx]
    pt = _table(ws, r + 1, ["Date", f"P10 ({sym})", f"Median ({sym})", f"P90 ({sym})",
                            f"Invested capital ({sym})"], prow,
                [FMT_DATE, money0, money0, money0, money0], autofilter=False)
    ws.add_chart(_date_line_chart(ws, r + 1, len(prow), 1, [2, 3, 4, 5],
                                  f"Projected NAV — {n_sims:,} simulations ({sym})",
                                  f"NAV ({sym})", "#,##0",
                                  [NEG_SOFT, GOLD, POS, NAVY],
                                  dashes=[None, None, None, "dash"],
                                  widths=[1.5, 2.5, 1.5, 1.75],
                                  span_days=horizon * 365), "H4")
    # Terminal distribution
    counts, edges = np.histogram(terminal, bins=30)
    r = pt["next"] + 1
    _section(ws, r, 1, "Terminal NAV distribution", 3)
    hrows = [[f"{(edges[i] + edges[i + 1]) / 2:,.0f}", (edges[i] + edges[i + 1]) / 2,
              int(counts[i])] for i in range(len(counts))]
    _table(ws, r + 1, ["Bin (mid)", f"NAV ({sym})", "Paths"], hrows,
           [None, money0, FMT_INT], autofilter=False)
    ch = _bar_chart(ws, r + 1, len(hrows), 1, 3, f"Terminal NAV after {horizon} year(s)",
                    "Number of paths", "#,##0", sign_colors=False, color=GOLD,
                    x_title=f"Terminal NAV ({sym})")
    ch.gapWidth = 10
    ws.add_chart(ch, "H28")
    _widths(ws, {"A": 34, "B": 18, "C": 18, "D": 18, "E": 22})
